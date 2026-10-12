"""The four format and parser tools as Plan O's standalone specs (M240): shape, safety, and what they return."""

from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path
from typing import Any

import pandas as pd
import pytest
from openpyxl import Workbook  # type: ignore[import-untyped]
from pydantic import BaseModel, ValidationError

from engine.doctor.tools_formats import (
    FORMAT_TOOLS,
    CheckParserLossInput,
    FindExcelSerialDatesInput,
    FormatToolSpec,
    ReadOtherFormatInput,
    ReadWorkbookSheetsInput,
    register,
)
from engine.storage import LocalStorage

pytestmark = pytest.mark.filterwarnings("ignore::pandas.errors.ParserWarning")

SECRET = "jane.roe@example.test"


class FakeContext:
    """A case with files by id: `received_file` and `derived_key` stay inside one client's folder."""

    def __init__(
        self, storage: LocalStorage, files: dict[str, tuple[str, str]], frames: dict[str, Any]
    ) -> None:
        self.storage = storage
        self._files = files
        self._frames = frames

    def received_file(self, file_id: str) -> tuple[str, str]:
        return self._files[file_id]

    def derived_key(self, file_id: str, name: str) -> str:
        return f"cases/c1/{file_id}/{name}"

    def frame(self, file_id: str) -> pd.DataFrame:
        return self._frames[file_id]  # type: ignore[no-any-return]


@pytest.fixture
def storage(tmp_path: Path) -> LocalStorage:
    return LocalStorage(tmp_path / "store")


def spec(name: str) -> FormatToolSpec:
    return next(tool for tool in FORMAT_TOOLS if tool.name == name)


def call(name: str, given: dict[str, Any], ctx: FakeContext) -> dict[str, Any]:
    tool = spec(name)
    out = tool.run(tool.input_model.model_validate(given), ctx)  # type: ignore[arg-type]
    assert isinstance(out, tool.output_model)
    return tool.output_model.model_validate(out.model_dump()).model_dump(mode="json")


# ---------------------------------------------------------------------------
# The specs
# ---------------------------------------------------------------------------
def test_the_four_tools_are_read_rules_tools_a_viewer_may_call() -> None:
    assert [tool.name for tool in FORMAT_TOOLS] == [
        "read_other_format",
        "read_workbook_sheets",
        "check_parser_loss",
        "find_excel_serial_dates",
    ]
    assert {(t.kind, t.mode, t.role) for t in FORMAT_TOOLS} == {("read", "rules", "viewer")}
    assert {t.name: t.exposure for t in FORMAT_TOOLS} == {
        "read_other_format": "session-only",
        "read_workbook_sheets": "session-only",
        "check_parser_loss": "aggregate",
        "find_excel_serial_dates": "session-only",
    }


def test_every_model_refuses_unknown_fields_and_no_tool_takes_a_storage_key() -> None:
    for tool in FORMAT_TOOLS:
        assert tool.input_model.model_config.get("extra") == "forbid", tool.name
        assert tool.output_model.model_config.get("extra") == "forbid", tool.name
        assert set(tool.input_model.model_fields) <= {"file_id", "sheets"}, tool.name
        with pytest.raises(ValidationError):
            tool.input_model.model_validate({"file_id": "f1", "key": "other-client/secret.csv"})
        assert tool.description and tool.wraps.startswith("engine.")


def test_register_hands_each_spec_to_the_registry_once() -> None:
    class Registry:
        def __init__(self) -> None:
            self.names: list[str] = []

        def register(self, tool: FormatToolSpec) -> None:
            self.names.append(tool.name)

    registry = Registry()
    register(registry)
    assert registry.names == [tool.name for tool in FORMAT_TOOLS]


# ---------------------------------------------------------------------------
# What they return
# ---------------------------------------------------------------------------
def test_read_other_format_reads_and_reports_without_a_cell(storage: LocalStorage) -> None:
    storage.write_bytes("in/f1.json", json.dumps([{"id": "0007", "email": SECRET}, {"id": "0008"}]).encode())
    ctx = FakeContext(storage, {"f1": ("in/f1.json", "customers.json")}, {})
    out = call("read_other_format", {"file_id": "f1"}, ctx)
    assert out["status"] == "read" and out["rows"] == 2 and out["columns"] == ["id", "email"]
    assert out["inner_format"] == "json" and out["container"] == "none"
    assert storage.exists("cases/c1/f1/source.csv")
    assert SECRET not in json.dumps(out) and "0007" not in json.dumps(out)


def test_read_other_format_turns_a_refusal_into_a_result_the_agent_can_explain(storage: LocalStorage) -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("a.csv", "x\n1\n")
        archive.writestr("b.csv", "x\n2\n")
    storage.write_bytes("in/f2.zip", buffer.getvalue())
    ctx = FakeContext(storage, {"f2": ("in/f2.zip", "two.zip")}, {})
    out = call("read_other_format", {"file_id": "f2"}, ctx)
    assert out["status"] == "refused" and out["code"] == "UPLOAD_ARCHIVE_NOT_ONE_TABLE"
    assert "holds 2 tables" in out["message"] and out["rows"] is None
    assert not storage.exists("cases/c1/f2/source.csv")


def test_read_workbook_sheets_writes_one_file_per_table_sheet(storage: LocalStorage) -> None:
    book = Workbook()
    book.remove(book.active)
    customers = book.create_sheet("Customers")
    customers.append(["id", "city"])
    customers.append(["0001", "Pune"])
    customers.append(["0002", "Delhi"])
    bills = book.create_sheet("Bills")
    bills.append(["id", "amount"])
    bills.append(["0001", 10])
    bills.append(["0002", 20])
    book.create_sheet("Cover").append(["Prepared by finance"])
    buffer = io.BytesIO()
    book.save(buffer)
    storage.write_bytes("in/f3.xlsx", buffer.getvalue())
    ctx = FakeContext(storage, {"f3": ("in/f3.xlsx", "crm.xlsx")}, {})
    out = call("read_workbook_sheets", {"file_id": "f3"}, ctx)
    assert out["status"] == "read"
    assert [(f["sheet"], f["rows"], f["file_position"]) for f in out["files"]] == [
        ("Customers", 2, 1),
        ("Bills", 2, 2),
    ]
    assert out["left_out"] == [{"sheet": "Cover", "reason": "no_table"}]
    assert storage.exists("cases/c1/f3/sheet-1.csv") and storage.exists("cases/c1/f3/sheet-2.csv")
    refused = call("read_workbook_sheets", {"file_id": "f3", "sheets": ["Orders"]}, ctx)
    assert refused["status"] == "refused" and refused["code"] == "UPLOAD_SHEET_NOT_FOUND"


def test_check_parser_loss_returns_the_probe_result_with_lines_only(storage: LocalStorage) -> None:
    rows = [f"C{i},Gold,{i}" for i in range(1, 11)]
    rows[4] = f"C5,{SECRET},5,EXTRA"
    storage.write_bytes("in/f4.csv", ("id,plan,spend\n" + "\n".join(rows) + "\n").encode())
    ctx = FakeContext(storage, {"f4": ("in/f4.csv", "customers.csv")}, {})
    out = call("check_parser_loss", {"file_id": "f4"}, ctx)
    assert (out["verdict"], out["code"], out["lines"], out["rows_dropped"]) == (
        "not_ready",
        "X_PARSER_ROW_LOSS",
        [6],
        1,
    )
    assert SECRET not in json.dumps(out) and "C5" not in json.dumps(out)


def test_find_excel_serial_dates_names_columns_and_the_step_parameters(storage: LocalStorage) -> None:
    frame = pd.DataFrame(
        {
            "id": list("abcdef"),
            "signup_date": [44927, 45000, 45292, 44000, 43000, 45100],
            "amount": [44927] * 6,
        }
    )
    ctx = FakeContext(storage, {}, {"f5": frame})
    out = call("find_excel_serial_dates", {"file_id": "f5"}, ctx)
    assert [c["column"] for c in out["columns"]] == ["signup_date"]
    column = out["columns"][0]
    assert (column["excel_serial"], column["date_1904"], column["convertible"], column["failed"]) == (
        True,
        False,
        6,
        0,
    )


def test_the_input_models_are_the_models_the_specs_declare() -> None:
    models: dict[str, type[BaseModel]] = {
        "read_other_format": ReadOtherFormatInput,
        "read_workbook_sheets": ReadWorkbookSheetsInput,
        "check_parser_loss": CheckParserLossInput,
        "find_excel_serial_dates": FindExcelSerialDatesInput,
    }
    assert {tool.name: tool.input_model for tool in FORMAT_TOOLS} == models


def test_a_result_goes_through_the_egress_gate_with_a_name_from_inside_the_file_masked(
    storage: LocalStorage,
) -> None:
    """The one text in a result that came from the person's file is an archive member's name; the gate masks it."""
    from engine.agent.egress import prepare

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(f"{SECRET}.csv", "x\n1\n")
        archive.writestr("b.csv", "x\n2\n")
    storage.write_bytes("in/f6.zip", buffer.getvalue())
    ctx = FakeContext(storage, {"f6": ("in/f6.zip", "two.zip")}, {})
    out = call("read_other_format", {"file_id": "f6"}, ctx)
    assert SECRET in out["message"]  # the tool's own result names the file, as the person sees it
    safe = prepare(out, ctx_columns=[], personal_columns=[])
    assert SECRET not in json.dumps(safe) and "example.test" not in json.dumps(safe)
    assert safe["code"] == "UPLOAD_ARCHIVE_NOT_ONE_TABLE"
