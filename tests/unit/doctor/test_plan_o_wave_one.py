"""Plan O wave 1 integration (DEC-1802 (t)): the codes, the tools of M234/M237/M238/M240/M241 in M232's registry,
M240's file context over M231's case, and M237's phone rule through M241's.

What is pinned here:

* every code the wave's modules raise is in `PLAN_O_CODES` (and so has a help entry, `test_help.py`);
* the wave's tools are registered by their catalogue rows, for the rules and the screen only - never the AI, so
  the chat's tool list is the same with `data_doctor.unified` on - and a row-level tool needs an Analyst;
* a module's own coded refusal reaches the caller with its code;
* M240's tools run on M231's case: a file outside the case is refused, and a workbook's two sheets become two
  files of one case (M240's acceptance, through the real case store);
* M237's `normalise_key` writes an Indian mobile number with M241's `to_e164_india`.
"""

from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Any

import pandas as pd
import pytest
from openpyxl import Workbook  # type: ignore[import-untyped]

from engine.access.roles import Principal, Role
from engine.agent.key_steps import KeyRule, normalise_key_value
from engine.agent.tools import AgentContext
from engine.doctor.build_driver import BUILD_CODES
from engine.doctor.case import CASE_CODES, CaseFile, CaseStore
from engine.doctor.case_files import DERIVED_DIR, CaseFiles
from engine.doctor.catalogue import row
from engine.doctor.codes import PLAN_O_CODES
from engine.doctor.diagnose import LoadedFile
from engine.doctor.india_ids import to_e164_india
from engine.doctor.join_files import JOIN_CODES
from engine.doctor.parser_loss import PARSER_LOSS_CODES
from engine.doctor.readiness_contracts import READINESS_CODES
from engine.doctor.registry import DataToolContext, DataToolError, DataToolKind, Exposure, data_registry
from engine.doctor.stack_files import STACK_CODES
from engine.doctor.tools_formats import FORMAT_TOOLS
from engine.doctor.tools_rows import ROW_TOOL_CODES
from engine.doctor.tools_wave_one import UNCATALOGUED, WAVE_ONE_TOOLS, _FilesView
from engine.pilot.help import known_codes
from engine.storage import LocalStorage, upload_key
from engine.utils.time import utc_now
from tests.unit.agent.egress_canary import canary_context

ANALYST = Principal(user_id="u-analyst", username="analyst", roles=frozenset({Role.ANALYST}), kind="user")
VIEWER = Principal(user_id="u-viewer", username="viewer", roles=frozenset({Role.VIEWER}), kind="user")


# ---------------------------------------------------------------------------
# Codes
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    "codes",
    [CASE_CODES, READINESS_CODES, PARSER_LOSS_CODES, ROW_TOOL_CODES, JOIN_CODES, STACK_CODES, BUILD_CODES],
    ids=["M231", "M234", "M240", "M238", "M237-join", "M237-stack", "M237-build"],
)
def test_every_code_the_wave_raises_is_a_plan_o_code_with_help(codes: frozenset[str]) -> None:
    assert codes <= PLAN_O_CODES
    assert codes <= known_codes()


def test_the_exact_text_code_m237_raises_is_a_plan_o_code() -> None:
    assert "APPROVAL_SENTENCE_CHANGED" in PLAN_O_CODES


# ---------------------------------------------------------------------------
# The registry
# ---------------------------------------------------------------------------
def test_the_wave_registers_every_catalogued_tool_by_its_row() -> None:
    registry = data_registry()
    for name in WAVE_ONE_TOOLS:
        tool = registry.get(name)
        entry = row(name)
        assert entry is not None and entry.number == tool.number, name
        assert entry.kind == tool.kind.value and entry.exposure == tool.exposure.value, name
        if tool.kind is DataToolKind.PROPOSE:
            assert tool.policy_kind == f"data.{name}", name
    for name in UNCATALOGUED:
        assert row(name) is None and name not in registry.names()
    assert {spec.name for spec in FORMAT_TOOLS} - UNCATALOGUED <= set(WAVE_ONE_TOOLS)
    # Plan O M236 (DEC-1806): M240's two tools have rows #82 and #83 now, so every format tool is registered.
    assert not UNCATALOGUED and {spec.name for spec in FORMAT_TOOLS} <= set(registry.names())


def test_the_wave_serves_the_rules_and_the_screen_never_the_ai() -> None:
    for name in WAVE_ONE_TOOLS:
        assert data_registry().get(name).callers == frozenset({"rules", "person"}), name


def test_a_row_level_wave_tool_needs_an_analyst_and_the_readiness_verdict_a_viewer() -> None:
    for name in WAVE_ONE_TOOLS:
        tool = data_registry().get(name)
        expected = Role.VIEWER if tool.exposure is Exposure.AGGREGATE else Role.ANALYST
        assert tool.route_policy.role is expected, name
    assert data_registry().get("check_readiness").exposure is Exposure.AGGREGATE


@pytest.fixture(scope="module")
def agent() -> AgentContext:
    return canary_context(rows=300)


def test_the_chat_lists_the_same_tools_with_the_wave_registered(agent: AgentContext) -> None:
    ctx = DataToolContext(agent=agent, principal=ANALYST, caller="ai", primary_key="customer_id")
    listed = {t.name for t in data_registry().listed_for(ctx, DataToolKind.READ)}
    assert listed.isdisjoint(WAVE_ONE_TOOLS)


def test_a_viewer_cannot_call_a_row_level_wave_tool(agent: AgentContext) -> None:
    ctx = DataToolContext(agent=agent, principal=VIEWER, caller="person")
    with pytest.raises(DataToolError) as refused:
        data_registry().call("impossible_values", {}, ctx)
    assert refused.value.code == "TOOL_ROLE_REQUIRED" and refused.value.status == 403


def test_a_wave_tool_whose_need_is_missing_is_not_available(agent: AgentContext) -> None:
    ctx = DataToolContext(agent=agent, principal=ANALYST, caller="person")
    for name in ("check_parser_loss", "check_readiness", "join_files"):
        assert data_registry().find(name, ctx) is None, name


def test_a_modules_coded_refusal_reaches_the_caller_with_its_code(agent: AgentContext) -> None:
    files = (LoadedFile(upload_id="u1", file_name="customers.csv", frame=agent.frame, profile=agent.profile),)
    ctx = DataToolContext(agent=agent, principal=ANALYST, caller="person", files=files)
    with pytest.raises(DataToolError) as refused:
        data_registry().call(
            "join_files",
            {"base_upload_id": "u1", "base_key_column": "no_such_column", "upload_ids": ["u1"]},
            ctx,
        )
    assert refused.value.code == "JOIN_COLUMN_MISSING" and refused.value.status == 422


def test_m238s_tools_read_the_decided_id_column(agent: AgentContext) -> None:
    ctx = DataToolContext(agent=agent, principal=ANALYST, caller="person", primary_key="customer_id")
    with pytest.raises(DataToolError) as refused:
        data_registry().call("detect_units", {"column": "customer_id"}, ctx)
    assert refused.value.code == "DOCTOR_COLUMN_PROTECTED"


# ---------------------------------------------------------------------------
# M240's tools on M231's case
# ---------------------------------------------------------------------------
def _upload(storage: LocalStorage, upload_id: str, name: str, data: bytes, stored: str) -> None:
    storage.write_bytes(upload_key(upload_id, stored), data)
    storage.write_text(
        upload_key(upload_id, "upload.json"),
        json.dumps({"upload_id": upload_id, "file_name": name, "source_key": upload_key(upload_id, stored)}),
    )


def _workbook() -> bytes:
    book = Workbook()
    book.remove(book.active)
    customers = book.create_sheet("Customers")
    customers.append(["id", "city"])
    customers.append(["0001", "Pune"])
    customers.append(["0002", "Delhi"])
    bills = book.create_sheet("Bills")
    bills.append(["id", "amount"])
    bills.append(["0001", 10])
    bills.append(["0002", 20])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def case_files(tmp_path: Path) -> CaseFiles:
    storage = LocalStorage(tmp_path / "store")
    _upload(storage, "u-book", "crm.xlsx", _workbook(), "original.xlsx")
    _upload(storage, "u-csv", "bills.csv", b"id,amount\n1,10\n2,20,30\n3,40\n", "source.csv")
    _upload(storage, "u-other", "other.csv", b"id\n1\n", "source.csv")
    store = CaseStore(storage)
    files = [
        CaseFile(upload_id=upload_id, file_name=name, added_at=utc_now(), added_by=ANALYST.user_id)
        for upload_id, name in (("u-book", "crm.xlsx"), ("u-csv", "bills.csv"))
    ]
    case = store.create(client_id="default", created_by=ANALYST.user_id, files=files)
    made: list[str] = []

    def make_upload(csv_key: str, file_name: str) -> str:
        upload_id = f"u-sheet-{len(made) + 1}"
        made.append(upload_id)
        _upload(storage, upload_id, file_name, storage.read_bytes(csv_key), "source.csv")
        return upload_id

    return CaseFiles(store, "default", case.case_id, ANALYST.user_id, make_upload=make_upload)


def _call(name: str, given: dict[str, Any], case_files: CaseFiles) -> dict[str, Any]:
    ctx = DataToolContext(principal=ANALYST, caller="person", case_files=case_files)
    return data_registry().call(name, given, ctx).output


def test_a_file_outside_the_case_is_refused(case_files: CaseFiles) -> None:
    with pytest.raises(DataToolError) as refused:
        _call("check_parser_loss", {"file_id": "u-other"}, case_files)
    assert refused.value.code == "TOOL_INPUT_INVALID"


def test_the_parser_check_reads_the_case_files_stored_csv(case_files: CaseFiles) -> None:
    out = _call("check_parser_loss", {"file_id": "u-csv"}, case_files)
    assert out["status"] == "checked" and out["report"]["rows_dropped"] == 1


def test_derived_files_live_inside_their_uploads_directory(case_files: CaseFiles) -> None:
    assert case_files.derived_key("u-csv", "source.csv") == f"uploads/u-csv/{DERIVED_DIR}/source.csv"
    assert case_files.readable_key("u-csv") == "uploads/u-csv/source.csv"
    assert case_files.received_file("u-book") == ("uploads/u-book/original.xlsx", "crm.xlsx")
    with pytest.raises(DataToolError):
        case_files.derived_key("u-csv", "../escape.csv")


def test_a_workbooks_two_sheets_become_two_files_of_one_case(case_files: CaseFiles) -> None:
    """M240's acceptance through M231's case store: the tool runs on the spec, the context is the real case."""
    spec = next(tool for tool in FORMAT_TOOLS if tool.name == "read_workbook_sheets")
    out = spec.run(spec.input_model.model_validate({"file_id": "u-book"}), case_files)
    case = case_files.store.get(case_files.client_id, case_files.case_id)
    added = [file for file in case.files if file.original_upload_id == "u-book"]
    assert [file.file_name for file in added] == ["crm - Customers.csv", "crm - Bills.csv"]
    assert [file.upload_id for file in added] == [sheet.file_id for sheet in out.files]  # type: ignore[attr-defined]
    frame = case_files.frame(added[0].upload_id)
    assert list(frame.columns) == ["id", "city"] and len(frame) == 2


def test_without_a_way_to_make_an_upload_nothing_is_added(tmp_path: Path, case_files: CaseFiles) -> None:
    from dataclasses import replace

    bare = replace(case_files, make_upload=None)
    key = bare.derived_key("u-book", "sheet-1.csv")
    bare.storage.write_bytes(key, b"id\n1\n")
    with pytest.raises(DataToolError):
        bare.add_file("u-book", key, "Customers")
    assert len(bare.store.get(bare.client_id, bare.case_id).files) == 2


# ---------------------------------------------------------------------------
# M237's phone rule is M241's
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    "written",
    ["98765 43210", "+91 98765 43210", "098765-43210", "0091 9876543210", "(+91) 98765 43210", "98765/43210"],
)
def test_normalise_key_writes_an_indian_mobile_as_to_e164_india_does(written: str) -> None:
    assert to_e164_india(written) == "+919876543210"
    assert normalise_key_value(written, KeyRule(phone=True)) == ("+919876543210", True)


def test_normalise_key_keeps_its_wider_reading_where_m241_reads_no_mobile() -> None:
    assert normalise_key_value("+1 415 555 0100", KeyRule(phone=True)) == ("+14155550100", True)
    assert normalise_key_value("0221234567", KeyRule(phone=True)) == ("+910221234567", True)
    assert normalise_key_value("not a phone", KeyRule(phone=True)) == ("not a phone", False)


def test_the_frame_helper_stays_a_dataframe(case_files: CaseFiles) -> None:
    assert isinstance(case_files.frame("u-csv"), pd.DataFrame)


def test_a_file_with_no_name_is_named_by_its_upload_id() -> None:
    """Wave-2 integration, as fix-build-sentence did in `build_driver.plan_build`: a file tool's sentence never
    names an empty file ("customers.csv and ."); it names the upload id."""
    frame = pd.DataFrame({"customer_id": ["C-1"]})
    files = tuple(
        LoadedFile(upload_id=upload_id, file_name=name, frame=frame, profile=None)  # type: ignore[arg-type]
        for upload_id, name in (("up-blank", ""), ("up-spaces", "  "), ("up-named", "bills.csv"))
    )
    view = _FilesView(DataToolContext(files=files))
    assert view.file_name("up-blank") == "up-blank"
    assert view.file_name("up-spaces") == "up-spaces"
    assert view.file_name("up-named") == "bills.csv"
