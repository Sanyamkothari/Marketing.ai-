"""Each sheet of a workbook as its own file (Plan O M240, DEC-1810 (e))."""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from openpyxl import Workbook  # type: ignore[import-untyped]

from engine.intake.sheets import MAX_SHEETS, split_workbook
from engine.stages.ingest import IngestError, read_upload
from engine.storage import LocalStorage


@pytest.fixture
def storage(tmp_path: Path) -> LocalStorage:
    return LocalStorage(tmp_path / "store")


def workbook(storage: LocalStorage, build: object, key: str = "u/crm_export.xlsx") -> str:
    book = Workbook()
    book.remove(book.active)
    build(book)  # type: ignore[operator]
    buffer = io.BytesIO()
    book.save(buffer)
    storage.write_bytes(key, buffer.getvalue())
    return key


def customers_and_bills(book: Workbook) -> None:
    customers = book.create_sheet("Customers")
    customers.append(["customer_id", "city", "joined"])
    for index in range(1, 6):
        customers.append([f"{index:07d}", ("Pune", "Delhi")[index % 2], f"2024-0{index}-01"])
    bills = book.create_sheet("Bills")
    bills.append(["Monthly bills", None, None])  # a title row above the table
    bills.append(["customer_id", "month", "amount"])
    for index in range(1, 9):
        bills.append([f"{(index % 5) + 1:07d}", f"2024-{index:02d}", 100 * index])
    bills.append(["Total", None, 3600])


def keys(position: int, sheet: str) -> str:
    return f"files/{position}/source.csv"


def test_a_workbook_with_a_customer_sheet_and_a_bills_sheet_becomes_two_files(storage: LocalStorage) -> None:
    key = workbook(storage, customers_and_bills)
    split = split_workbook(storage, received_key=key, csv_key_for=keys, file_name="crm_export.xlsx")
    assert [(f.sheet, f.rows, f.csv_key) for f in split.files] == [
        ("Customers", 5, "files/1/source.csv"),
        ("Bills", 8, "files/2/source.csv"),
    ]
    assert [f.label for f in split.files] == ["crm_export - Customers", "crm_export - Bills"]
    assert split.left_out == () and {info.name for info in split.sheets} == {"Customers", "Bills"}
    customers = read_upload(storage, "files/1/source.csv", keep_leading_zeros=True).frame
    bills = read_upload(storage, "files/2/source.csv", keep_leading_zeros=True).frame
    assert list(customers.columns) == ["customer_id", "city", "joined"]
    assert list(bills.columns) == [
        "customer_id",
        "month",
        "amount",
    ]  # title row and total row left out by intake
    assert customers["customer_id"].tolist()[0] == "0000001"
    assert set(bills["customer_id"]) <= set(customers["customer_id"])  # the two files join
    assert all(f.report.original_key == key and f.report.sheet == f.sheet for f in split.files)


def test_a_hidden_sheet_and_a_sheet_with_no_table_are_left_out_and_said_so(storage: LocalStorage) -> None:
    def build(book: Workbook) -> None:
        customers_and_bills(book)
        cover = book.create_sheet("Cover")
        cover.append(["Prepared by finance"])
        secret = book.create_sheet("Archive")
        secret.append(["a", "b"])
        secret.append([1, 2])
        secret.sheet_state = "hidden"

    key = workbook(storage, build)
    split = split_workbook(storage, received_key=key, csv_key_for=keys)
    assert [f.sheet for f in split.files] == ["Customers", "Bills"]
    assert {(item.sheet, item.reason) for item in split.left_out} == {
        ("Cover", "no_table"),
        ("Archive", "hidden"),
    }


def test_sheets_can_be_named_and_a_hidden_one_is_read_when_named(storage: LocalStorage) -> None:
    def build(book: Workbook) -> None:
        customers_and_bills(book)
        secret = book.create_sheet("Archive")
        secret.append(["a", "b"])
        secret.append([1, 2])
        secret.sheet_state = "hidden"

    key = workbook(storage, build)
    split = split_workbook(storage, received_key=key, csv_key_for=keys, sheets=("archive", "Bills "))
    assert [f.sheet for f in split.files] == ["Bills", "Archive"]
    assert split.left_out == ()


def test_a_sheet_that_does_not_exist_is_refused_listing_the_ones_that_do(storage: LocalStorage) -> None:
    key = workbook(storage, customers_and_bills)
    with pytest.raises(IngestError) as caught:
        split_workbook(storage, received_key=key, csv_key_for=keys, sheets=("Orders",))
    assert caught.value.code == "UPLOAD_SHEET_NOT_FOUND"
    assert '"Customers"' in caught.value.message and '"Bills"' in caught.value.message
    assert not storage.list_keys("files/")


def test_a_workbook_with_no_table_is_refused(storage: LocalStorage) -> None:
    def build(book: Workbook) -> None:
        book.create_sheet("Notes").append(["just a note"])

    key = workbook(storage, build)
    with pytest.raises(IngestError) as caught:
        split_workbook(storage, received_key=key, csv_key_for=keys)
    assert caught.value.code == "UPLOAD_NO_ROWS"


def test_not_a_workbook_is_refused_as_intake_refuses_it(storage: LocalStorage) -> None:
    storage.write_bytes("u/bad.xlsx", b"not a workbook")
    with pytest.raises(IngestError) as caught:
        split_workbook(storage, received_key="u/bad.xlsx", csv_key_for=keys)
    assert caught.value.code == "UPLOAD_WORKBOOK_UNREADABLE"


def test_a_workbook_with_more_sheets_than_the_limit_writes_the_limit_and_names_the_rest(
    storage: LocalStorage,
) -> None:
    def build(book: Workbook) -> None:
        for index in range(MAX_SHEETS + 3):
            sheet = book.create_sheet(f"S{index}")
            sheet.append(["a", "b"])
            sheet.append([1, 2])
            sheet.append([3, 4])

    key = workbook(storage, build)
    split = split_workbook(storage, received_key=key, csv_key_for=keys)
    assert len(split.files) == MAX_SHEETS
    assert [item.reason for item in split.left_out] == ["over_limit"] * 3 and len(
        storage.list_keys("files/")
    ) == 20
