"""`intake_file` on stored files (Plan L M150, DEC-1501, DEC-1503): the edges the bench does not cover."""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Callable
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook  # type: ignore[import-untyped]

from engine.intake.config import IntakeConfig
from engine.intake.run import IntakeResult, intake_file, received_format_for
from engine.stages.ingest import IngestError, read_upload
from engine.storage import LocalStorage

CONFIG = IntakeConfig()


@pytest.fixture
def storage(tmp_path: Path) -> LocalStorage:
    return LocalStorage(tmp_path / "store")


def csv_intake(storage: LocalStorage, data: bytes, **kwargs: object) -> IntakeResult:
    storage.write_bytes("u/source.csv", data)
    return intake_file(
        storage,
        received_key="u/source.csv",
        received_format="csv",
        csv_key="u/source.csv",
        original_key="u/original.csv",
        config=kwargs.pop("config", CONFIG),  # type: ignore[arg-type]
        **kwargs,  # type: ignore[arg-type]
    )


def book_intake(storage: LocalStorage, book: Workbook, **kwargs: object) -> IntakeResult:
    buffer = io.BytesIO()
    book.save(buffer)
    storage.write_bytes("u/original.xlsx", buffer.getvalue())
    return intake_file(
        storage,
        received_key="u/original.xlsx",
        received_format="xlsx",
        csv_key="u/source.csv",
        original_key="u/original.xlsx",
        config=kwargs.pop("config", CONFIG),  # type: ignore[arg-type]
        **kwargs,  # type: ignore[arg-type]
    )


MESSY = b"Customer report,,\nid,plan,spend\nC1,Gold,10\nC2,Silver,-\nC3,Gold,30\nGrand Total,,40\n"


# ---------------------------------------------------------------------------
# Formats
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    ("name", "expected"),
    [("a.csv", "csv"), ("A.TXT", "csv"), ("a.parquet", "parquet"), ("a.xlsx", "xlsx"), ("a.XLSM", "xlsx")],
)
def test_received_formats(name: str, expected: str) -> None:
    assert received_format_for(name) == expected


@pytest.mark.parametrize("name", ["old.xls", "binary.xlsb"])
def test_an_old_excel_format_is_refused_with_what_to_do(name: str) -> None:
    with pytest.raises(IngestError) as caught:
        received_format_for(name)
    assert caught.value.code == "UPLOAD_UNSUPPORTED_FORMAT"
    assert "save it as .xlsx or CSV" in caught.value.message


def test_another_format_names_the_three_that_are_read() -> None:
    with pytest.raises(IngestError, match=r"CSV, Excel \(\.xlsx\) and Parquet"):
        received_format_for("notes.docx")


def test_parquet_is_read_as_it_is(storage: LocalStorage) -> None:
    result = intake_file(
        storage,
        received_key="u/source.parquet",
        received_format="parquet",
        csv_key="u/source.csv",
        original_key="u/original.parquet",
        config=CONFIG,
    )
    assert (result.source_key, result.file_format, result.report.applied) == (
        "u/source.parquet",
        "parquet",
        False,
    )


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------
def test_a_messy_csv_is_rewritten_and_its_original_kept(storage: LocalStorage) -> None:
    result = csv_intake(storage, MESSY)
    assert result.report.applied
    assert storage.read_bytes("u/original.csv") == MESSY
    frame = read_upload(storage, "u/source.csv").frame
    assert list(frame.columns) == ["id", "plan", "spend"]
    assert frame["spend"].tolist()[:1] == [10] and pd.isna(frame["spend"].iloc[1])
    assert "Grand Total" not in set(frame["id"])


def test_turned_off_reads_the_csv_exactly_as_before(storage: LocalStorage) -> None:
    result = csv_intake(storage, MESSY, enabled=False)
    assert result.report.turned_off and not result.report.applied
    assert storage.read_bytes("u/source.csv") == MESSY
    assert not storage.exists("u/original.csv")
    assert result.report.notes == ("Read as it is: the intake rules were turned off for this file.",)


def test_a_semicolon_file_with_a_title_finds_its_delimiter(storage: LocalStorage) -> None:
    data = b"Export of 3 Oct\nid;plan;spend;region\nC1;Gold;10;W\nC2;Silver;20;E\n"
    result = csv_intake(storage, data)
    assert result.report.delimiter == ";"
    assert list(read_upload(storage, "u/source.csv").frame.columns) == ["id", "plan", "spend", "region"]


def test_a_column_with_no_name_but_values_is_kept_under_the_pandas_name(storage: LocalStorage) -> None:
    data = b"Title,,\nid,plan,spend\nC1,Gold,10\nC2,Silver,20,stray\n"
    csv_intake(storage, data)
    frame = read_upload(storage, "u/source.csv").frame
    assert list(frame.columns) == ["id", "plan", "spend", "Unnamed: 3"]


def test_cells_past_the_rows_looked_at_are_counted_not_silently_dropped(storage: LocalStorage) -> None:
    body = b"".join(b"C%d,Gold,%d\n" % (index, index) for index in range(150))
    data = b"Title,,\nid,plan,spend\n" + body + b"C150,Gold,1,stray\n"
    result = csv_intake(storage, data, config=IntakeConfig(sample_rows=100))
    assert len(read_upload(storage, "u/source.csv").frame) == 151
    assert result.report.cells_left_out == 1
    assert any("outside the table" in note for note in result.report.notes)


def test_a_codec_that_fails_late_is_replaced(storage: LocalStorage) -> None:
    """The head is plain ASCII; a Windows-1252 byte far past it makes UTF-8 fail mid-rewrite."""
    body = b"".join(b"C%d,Gold,%d\n" % (index, index) for index in range(8000))
    data = b"Customer report,,\nid,plan,spend\n" + body + b"C9,Caf\xe9,1\n"
    result = csv_intake(storage, data)
    assert result.report.applied
    frame = read_upload(storage, "u/source.csv").frame
    assert frame["plan"].iloc[-1] == "Café"


def test_a_footer_label_is_masked(storage: LocalStorage) -> None:
    data = b"id,plan,spend\nC1,Gold,10\nC2,Silver,20\nC3,Gold,30\nSent by: ops.lead@example.com,,\n"
    result = csv_intake(storage, data)
    assert result.report.footer_rows == 1
    assert "example.com" not in result.report.footer_labels[0]


def test_a_clean_file_with_only_a_trailing_comma_line_loses_the_empty_row(storage: LocalStorage) -> None:
    data = b"id,plan,spend\nC1,Gold,10\nC2,Silver,20\n,,\n"
    result = csv_intake(storage, data)
    assert result.report.empty_rows == 1
    assert len(read_upload(storage, "u/source.csv").frame) == 2


def test_an_empty_file_is_left_to_ingest(storage: LocalStorage) -> None:
    result = csv_intake(storage, b"")
    assert not result.report.applied


# ---------------------------------------------------------------------------
# Workbooks
# ---------------------------------------------------------------------------
def three_sheet_book() -> Workbook:
    book = Workbook()
    notes = book.active
    notes.title = "Notes"
    notes.append(["Exported from the CRM"])
    customers = book.create_sheet("Customers")
    customers.append(["id", "plan", "spend"])
    for index in range(10):
        customers.append([f"C{index}", "Gold", index * 10])
    summary = book.create_sheet("Summary")
    summary.append(["plan", "customers"])
    summary.append(["Gold", 10])
    hidden = book.create_sheet("Everything")
    hidden.sheet_state = "hidden"
    hidden.append(["id", "secret"])
    for index in range(50):
        hidden.append([f"C{index}", index])
    return book


def test_the_largest_visible_table_is_read_and_the_others_listed(storage: LocalStorage) -> None:
    result = book_intake(storage, three_sheet_book())
    report = result.report
    assert report.sheet == "Customers"
    assert [(sheet.name, sheet.used, sheet.hidden) for sheet in report.sheets] == [
        ("Notes", False, False),
        ("Customers", True, False),
        ("Summary", False, False),
        ("Everything", False, True),
    ]
    assert any('"Summary"' in note for note in report.notes)
    assert len(read_upload(storage, "u/source.csv").frame) == 10


def undimensioned_book(sizes: tuple[int, ...]) -> Workbook:
    """A workbook written as pandas and many exporters write one: no `<dimension>`, so no row count."""
    book = Workbook(write_only=True)
    for index, size in enumerate(sizes):
        sheet = book.create_sheet(f"Region {index}")
        sheet.append(["Region report"])
        sheet.append(["id", "spend"])
        for row in range(size):
            sheet.append([f"C{row}", row])
    return book


def test_without_a_row_count_the_biggest_sheet_is_still_read(storage: LocalStorage) -> None:
    report = book_intake(storage, undimensioned_book((30, 300, 120))).report
    assert [sheet.rows for sheet in report.sheets] == [None, None, None]
    assert report.sheet == "Region 1"
    assert report.rows_written == 300


def test_opening_a_workbook_without_a_row_count_does_not_read_its_rows() -> None:
    from engine.intake.excel import open_workbook

    buffer = io.BytesIO()
    undimensioned_book((500, 500)).save(buffer)
    damaged = io.BytesIO()
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(damaged, "w") as target:
        for member in source.infolist():
            data = source.read(member)
            if member.filename.startswith("xl/worksheets/sheet"):
                # The end of every sheet is broken: only a reader that parses every row can notice.
                data = data.replace(b"</sheetData>", b"<row><broken></sheetData>")
            target.writestr(member, data)
    book = open_workbook(io.BytesIO(damaged.getvalue()))
    assert [sheet.max_row for sheet in book.worksheets] == [None, None]


def test_a_declared_row_count_is_still_read(storage: LocalStorage) -> None:
    report = book_intake(storage, three_sheet_book()).report
    assert [sheet.rows for sheet in report.sheets] == [1, 11, 2, 51]


def test_a_named_sheet_is_read_hidden_or_not_and_case_is_forgiven(storage: LocalStorage) -> None:
    assert book_intake(storage, three_sheet_book(), sheet="summary").report.sheet == "Summary"
    assert book_intake(storage, three_sheet_book(), sheet="Everything").report.sheet == "Everything"


def test_a_sheet_that_does_not_exist_names_the_ones_that_do(storage: LocalStorage) -> None:
    with pytest.raises(IngestError) as caught:
        book_intake(storage, three_sheet_book(), sheet="Bills")
    assert caught.value.code == "UPLOAD_SHEET_NOT_FOUND"
    assert '"Customers"' in caught.value.message


def test_cells_become_text_as_a_csv_export_writes_them(storage: LocalStorage) -> None:
    from datetime import datetime

    book = Workbook()
    sheet = book.active
    sheet.append(["id", "whole", "fraction", "when", "at", "flag"])
    sheet.append(["C1", 3.0, 0.25, datetime(2024, 5, 1), datetime(2024, 5, 1, 9, 30), True])
    sheet.append(["C2", 4, 1.5, datetime(2024, 6, 2), datetime(2024, 6, 2, 18, 0), False])
    book_intake(storage, book)
    assert storage.read_text("u/source.csv").splitlines() == [
        "id,whole,fraction,when,at,flag",
        "C1,3,0.25,2024-05-01,2024-05-01 09:30:00,TRUE",
        "C2,4,1.5,2024-06-02,2024-06-02 18:00:00,FALSE",
    ]


def test_turned_off_a_workbook_is_converted_and_nothing_else(storage: LocalStorage) -> None:
    book = Workbook()
    sheet = book.active
    for row in (["Report"], ["id", "spend"], ["C1", "-"], ["Total", 1]):
        sheet.append(row)
    result = book_intake(storage, book, enabled=False)
    assert result.report.turned_off and result.report.header_row == 1
    assert storage.read_text("u/source.csv").splitlines() == ["Report,", "id,spend", "C1,-", "Total,1"]


def test_a_file_that_is_not_a_workbook_is_refused(storage: LocalStorage) -> None:
    storage.write_bytes("u/original.xlsx", b"id,plan\n1,Gold\n")
    with pytest.raises(IngestError) as caught:
        intake_file(
            storage,
            received_key="u/original.xlsx",
            received_format="xlsx",
            csv_key="u/source.csv",
            original_key="u/original.xlsx",
            config=CONFIG,
        )
    assert caught.value.code == "UPLOAD_WORKBOOK_UNREADABLE"


def test_a_workbook_that_unpacks_past_the_limit_is_refused_unopened(storage: LocalStorage) -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/sharedStrings.xml", b"0" * (3 * 1024 * 1024))
    storage.write_bytes("u/original.xlsx", buffer.getvalue())
    small = IntakeConfig(max_workbook_uncompressed_mb=1)
    with pytest.raises(IngestError) as caught:
        intake_file(
            storage,
            received_key="u/original.xlsx",
            received_format="xlsx",
            csv_key="u/source.csv",
            original_key="u/original.xlsx",
            config=small,
        )
    assert caught.value.code == "UPLOAD_WORKBOOK_TOO_LARGE"


def test_an_empty_workbook_says_so(storage: LocalStorage) -> None:
    with pytest.raises(IngestError) as caught:
        book_intake(storage, Workbook())
    assert caught.value.code == "UPLOAD_NO_ROWS"


# ---------------------------------------------------------------------------
# Hostile workbooks (Plan L M161)
# ---------------------------------------------------------------------------
def rewritten(book: Workbook, change: Callable[[bytes], bytes]) -> bytes:
    """The workbook saved, with every sheet's XML passed through `change`."""
    buffer = io.BytesIO()
    book.save(buffer)
    out = io.BytesIO()
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for member in source.infolist():
            data = source.read(member)
            if member.filename.startswith("xl/worksheets/sheet"):
                data = change(data)
            target.writestr(member, data)
    return out.getvalue()


def raw_intake(storage: LocalStorage, data: bytes, config: IntakeConfig = CONFIG) -> IntakeResult:
    storage.write_bytes("u/original.xlsx", data)
    return intake_file(
        storage,
        received_key="u/original.xlsx",
        received_format="xlsx",
        csv_key="u/source.csv",
        original_key="u/original.xlsx",
        config=config,
    )


def small_table() -> Workbook:
    book = Workbook()
    sheet = book.active
    sheet.append(["id", "plan", "spend"])
    for index in range(20):
        sheet.append([f"C{index}", "Gold", index])
    return book


def test_a_declared_size_is_not_trusted(storage: LocalStorage) -> None:
    """openpyxl pads every row to the declared size: A1:XFD1048576 would be rows of 16,384 cells."""
    data = rewritten(
        small_table(),
        lambda xml: re.sub(rb'<dimension ref="[^"]*"/>', b'<dimension ref="A1:XFD1048576"/>', xml),
    )
    assert b"XFD1048576" in zipfile.ZipFile(io.BytesIO(data)).read("xl/worksheets/sheet1.xml")
    result = raw_intake(storage, data)
    assert result.report.rows_written == 20
    assert list(read_upload(storage, "u/source.csv").frame.columns) == ["id", "plan", "spend"]
    assert result.report.sheets[0].rows == 1048576, "the declared count is still reported"


def test_values_far_apart_in_every_row_are_refused_not_read_for_hours(storage: LocalStorage) -> None:
    book = Workbook()
    sheet = book.active
    sheet.append(["id", "spend"])
    for row in range(2, 202):
        sheet.cell(row=row, column=1, value=f"C{row}")
        sheet.cell(row=row, column=16384, value=row)
    with pytest.raises(IngestError) as caught:
        book_intake(storage, book)
    assert caught.value.code == "UPLOAD_WORKBOOK_TOO_SPARSE"


def test_too_many_texts_are_refused_unopened(storage: LocalStorage) -> None:
    """The table of texts is held in memory whole, so it has its own, smaller limit."""
    buffer = io.BytesIO()
    small_table().save(buffer)
    out = io.BytesIO()
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for member in source.infolist():
            target.writestr(member, source.read(member))
        target.writestr("xl/sharedStrings.xml", b"<sst>" + b"<si><t>x</t></si>" * 150_000 + b"</sst>")
    with pytest.raises(IngestError) as caught:
        raw_intake(storage, out.getvalue(), IntakeConfig(max_shared_strings_mb=1))
    assert caught.value.code == "UPLOAD_WORKBOOK_TOO_LARGE"
    assert "texts" in caught.value.message


def test_a_sheet_damaged_half_way_is_a_clear_refusal(storage: LocalStorage) -> None:
    data = rewritten(small_table(), lambda xml: xml.replace(b'r="A15"', b'r="A!15"'))
    with pytest.raises(IngestError) as caught:
        raw_intake(storage, data)
    assert caught.value.code == "UPLOAD_WORKBOOK_UNREADABLE"
