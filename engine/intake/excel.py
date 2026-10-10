"""Reading an Excel workbook as rows of text (Plan L M150, DEC-1503).

`openpyxl` in read-only mode streams a sheet's XML, so memory stays bounded by the shared-strings
table and one row. Two guards run before a workbook is opened: it must be a zip archive (an old `.xls`
or a password-protected file is not, and is refused with what to do instead), and its members must
not unpack to more than `max_workbook_uncompressed_mb`, so a small upload cannot expand into gigabytes.
The XML itself is parsed by the standard library's expat, whose built-in amplification limits stop
entity-expansion attacks; external entities are never resolved.

Cells become text the way a "Save as CSV" would write them, so the same rules read both alike: a
whole number without `.0`, a date as `YYYY-MM-DD` (with the time when it has one), `TRUE` / `FALSE`, an
empty cell as nothing. Formulas give their cached value (`data_only=True`); a formula Excel never
calculated reads as empty.
"""

from __future__ import annotations

import math
import zipfile
from collections.abc import Iterator
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, BinaryIO, Final

from engine.intake.config import IntakeConfig
from engine.intake.grid import find_header, nonempty

__all__ = [
    "EXCEL_SUFFIXES",
    "LEGACY_EXCEL_SUFFIXES",
    "SHEET_SCAN_ROWS",
    "WorkbookError",
    "cell_text",
    "check_workbook",
    "open_workbook",
    "sheet_rows",
]

EXCEL_SUFFIXES: Final[frozenset[str]] = frozenset({".xlsx", ".xlsm"})
LEGACY_EXCEL_SUFFIXES: Final[frozenset[str]] = frozenset({".xls", ".xlsb"})

SHEET_SCAN_ROWS: Final[int] = 200
"""Rows of each sheet read to decide whether it holds a table."""

_MIDNIGHT: Final[time] = time(0, 0)


class WorkbookError(Exception):
    """A workbook that cannot be read; `code` is an upload error code."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


def check_workbook(path: Path | BinaryIO, config: IntakeConfig) -> None:
    """Refuse a file that is not an .xlsx archive, or that would unpack past the configured limit."""
    if not zipfile.is_zipfile(path):
        raise WorkbookError(
            "UPLOAD_WORKBOOK_UNREADABLE",
            "This file is not a readable .xlsx workbook. If it is password-protected, remove the password; "
            "otherwise open it in Excel, save it as .xlsx or CSV, and upload it again.",
        )
    limit = config.max_workbook_uncompressed_mb * 1024 * 1024
    try:
        with zipfile.ZipFile(path) as archive:
            unpacked = sum(member.file_size for member in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise WorkbookError(
            "UPLOAD_WORKBOOK_UNREADABLE", "This .xlsx file is damaged and cannot be opened."
        ) from exc
    if unpacked > limit:
        raise WorkbookError(
            "UPLOAD_WORKBOOK_TOO_LARGE",
            f"This workbook unpacks to more than {config.max_workbook_uncompressed_mb:,} MB. "
            "Save the sheet you need as CSV and upload that instead.",
        )


def open_workbook(path: Path | BinaryIO) -> Any:
    """The workbook in read-only mode with cached formula values (`openpyxl` ships no type stubs)."""
    from openpyxl import load_workbook  # type: ignore[import-untyped]

    try:
        return load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a dozen unrelated types for a damaged file
        raise WorkbookError(
            "UPLOAD_WORKBOOK_UNREADABLE",
            "This workbook could not be opened. Open it in Excel, save it as .xlsx or CSV, and upload it again.",
        ) from exc


def cell_text(value: Any) -> str:
    """A cell's value as a CSV export would write it."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        if value.is_integer() and abs(value) < 2**53:
            return str(int(value))
        return repr(value)
    if isinstance(value, datetime):
        if value.time() == _MIDNIGHT and value.tzinfo is None:
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date | time):
        return value.isoformat()
    if isinstance(value, timedelta):
        return repr(value.total_seconds())
    return str(value)


def sheet_rows(sheet: Any, *, max_rows: int | None = None) -> Iterator[list[str]]:
    """The sheet's rows as text, from its first row, every row padded only as far as Excel stored it."""
    for row in sheet.iter_rows(max_row=max_rows, values_only=True):
        yield [cell_text(value) for value in row]


def looks_like_data(rows: list[list[str]], config: IntakeConfig) -> bool:
    """True when the rows hold a row of at least two column names with a filled row under it."""
    filled = [row for row in rows if nonempty(row)]
    if len(filled) < 2:
        return False
    decision = find_header(rows, config)
    header = rows[decision.header_index] if decision.header_index < len(rows) else []
    if len(nonempty(header)) < 2:
        return False
    return any(nonempty(row) for row in rows[decision.data_start :])
