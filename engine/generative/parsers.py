"""Reading an uploaded document into ordered, headed sections.

A retrieved passage is only as useful as the heading a reader is shown beside it. `Citation.section`
is the line that tells somebody *where in the refund policy* an answer came from, and a citation
that said "policy_refunds.pdf, somewhere" would be a citation nobody could check. So the whole job
of this module is to recover structure - which lines are headings, which page they sit on, and
which prose belongs under each - before anything is chunked, embedded or quoted.

**Four formats, one shape, one table.** :data:`DOCUMENT_PARSERS` maps a
:class:`engine.config.DocumentType` to the reader for it, and :func:`parse` looks the extension up
rather than branching on it. A chain of `if suffix == ".pdf"` would have put the accepted set in
two places - here and in `generative.knowledge_base.accepted_types` - and the two would have
drifted the first time a format was added. The enum is the single vocabulary; a suffix that is not
one of its values is `DOCUMENT_TYPE_UNSUPPORTED` before a byte is read.

**Structure is recovered, never invented.** DOCX carries real heading styles and Markdown carries
real ATX headings, so those two readers are told the answer. Plain text and PDF are not: both have
been flattened to lines and the heading information that was in the layout is gone. What is left is
a shape - a short, capitalised line that does not finish a sentence, sitting after a line that did -
and that shape is what :func:`_heading_shaped` tests. Two deliberate loosenings, because the strict
readings find nothing on real documents:

* *Sentence case, not Title Case.* "Prepaid packs" and "How a refund is paid" are headings in every
  policy document ever written, and neither is Title Case. The test is therefore the first
  character, not every word. ALL CAPS passes the same test and needs no separate branch.
* *The line before matters more than the line after.* A heading is preceded by the end of a
  sentence, by a blank line, or by another heading. That one rule is what separates a heading from
  the last, short line of a wrapped paragraph, and it is also what keeps a table's rows out: a row
  follows another row, and a row finishes no sentence. The line after is only checked for existence,
  because a heading that introduces a table is followed by a table, not by prose, and requiring
  prose there lost every such heading.

A table's cells are joined with :data:`CELL_SEPARATOR` by all four readers, so the same table reads
the same way whichever format it arrived in - and a line carrying that separator is never read as a
heading, since a flattened PDF renders an em dash the same way and the two cannot be told apart
mid-document. The first line of a PDF is exempt: it is the title, and titles carry em dashes.

**Free text hides its PII mid-sentence.** The Phase 1 detectors in `engine.stages.ingest` match a
whole cell, which is the right test for a column and the wrong one for a complaint. The patterns
themselves are correct and are reused verbatim through :data:`FREE_TEXT_DETECTORS` with
`re.finditer`; re-writing them here would have created a second definition of "what a phone number
looks like" and guaranteed the two would disagree. :func:`redact` replaces what it finds with
`engine.stages.ingest.REDACTED` and returns the *kinds*, never a matched value - the same discipline
`detect_pii` keeps, for the same reason (plan section 13.7).

`pypdf`, `python-docx` and `markdown-it-py` are imported inside the reader bodies, never at module
level, so `import engine` stays fast and the engine's import graph loads no document library
(tests/integration/test_engine_imports.py).
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from types import MappingProxyType
from typing import TYPE_CHECKING, Final

from engine.config import DocumentType
from engine.stages.ingest import PII_DETECTORS, REDACTED
from engine.utils.logging import get_logger, log_failure

if TYPE_CHECKING:
    from collections.abc import Callable, Mapping, Sequence
    from pathlib import Path

    from markdown_it.token import Token

    from engine.stages.ingest import PiiDetector

__all__ = [
    "CELL_SEPARATOR",
    "DOCUMENT_EMPTY",
    "DOCUMENT_NO_HEADINGS",
    "DOCUMENT_PAGE_UNREADABLE",
    "DOCUMENT_PARSERS",
    "DOCUMENT_TYPE_UNSUPPORTED",
    "DOCUMENT_UNREADABLE",
    "FREE_TEXT_DETECTORS",
    "HEADING_MAX_CHARS",
    "HEADING_MAX_WORDS",
    "PARSE_ERRORS",
    "ParseError",
    "ParsedDocument",
    "Section",
    "find_pii",
    "parse",
    "parse_error",
    "redact",
]

_LOGGER = get_logger(__name__)


# ---------------------------------------------------------------------------
# 1. Module constants
# ---------------------------------------------------------------------------
CELL_SEPARATOR: Final[str] = " - "
"""How every reader renders one table row: the cells joined, never the pipes.

A pipe is Markdown punctuation, not something a reader says out loud, and an answer that quoted
`| Starter 149 | 149 | 24 days |` would be quoting the syntax rather than the tariff. The same
join is used by `tests/fixtures/make_docs`, so a table survives a round trip through any of the
four formats as the same sentence.
"""

HEADING_MAX_CHARS: Final[int] = 70
"""Longest line a flattened format may still read as a heading.

Comfortably above the longest heading in a policy document and comfortably below a wrapped line of
prose, which the PDF writer breaks in the nineties. A number in between would let the last line of
a paragraph pass for a heading.
"""

HEADING_MAX_WORDS: Final[int] = 12
"""Most words a heading may have. A short sentence is longer than this; a heading rarely is."""

_HEADING_FORBIDDEN_ENDINGS: Final[str] = ".,;"
"""A heading finishes nothing. A question mark is allowed - an FAQ's headings are all questions."""

_SENTENCE_ENDINGS: Final[tuple[str, ...]] = (".", "?", "!", ":")
"""What the line *before* a heading ends with. A colon counts: it introduces what follows."""

_WHITESPACE_RUN: Final[re.Pattern[str]] = re.compile(r"\s+")
"""Extraction from a PDF spaces glyphs by position, so a single space arrives as several."""

_DOCX_HEADING_STYLES: Final[tuple[str, ...]] = ("Heading", "Title")
"""Style-name prefixes that open a section. `Title` is what `add_heading(level=0)` writes."""

_MD_BODY_TOKENS: Final[frozenset[str]] = frozenset({"fence", "code_block", "html_block"})
"""Markdown tokens whose `content` is body text in its own right rather than an inline run."""

_MD_TEXT_CHILDREN: Final[frozenset[str]] = frozenset({"text", "code_inline"})
"""Inline children that carry characters. Emphasis markers carry none and are dropped."""

_MD_BREAK_CHILDREN: Final[frozenset[str]] = frozenset({"softbreak", "hardbreak"})
"""Inline children that carry a line break, which rejoins as one space: a paragraph is a paragraph."""

DOCUMENT_NO_HEADINGS: Final[str] = "DOCUMENT_NO_HEADINGS"
"""The parser found no heading, so the whole document is one section named after the file."""

DOCUMENT_PAGE_UNREADABLE: Final[str] = "DOCUMENT_PAGE_UNREADABLE"
"""One page of a PDF would not yield text. The rest of the document is still indexed."""


# ---------------------------------------------------------------------------
# 2. The error family
# ---------------------------------------------------------------------------
DOCUMENT_TYPE_UNSUPPORTED: Final[str] = "DOCUMENT_TYPE_UNSUPPORTED"
"""The extension is not one of `DocumentType`'s values, so no reader in the table owns it."""

DOCUMENT_UNREADABLE: Final[str] = "DOCUMENT_UNREADABLE"
"""The file is missing, or the library that owns the format refused to open it."""

DOCUMENT_EMPTY: Final[str] = "DOCUMENT_EMPTY"
"""The file opened and carried no text: a scan with no text layer, or an empty file."""


PARSE_ERRORS: Final[Mapping[str, tuple[str, str]]] = MappingProxyType(
    {
        DOCUMENT_TYPE_UNSUPPORTED: (
            "{name} is a {suffix} file, and the knowledge base reads only PDF, Word, Markdown "
            "and plain text.",
            "Save the document as a PDF or a Word file and upload it again.",
        ),
        DOCUMENT_UNREADABLE: (
            "{name} could not be opened, so nothing could be read from it.",
            "Open the file yourself to check it is complete, then upload it again.",
        ),
        DOCUMENT_EMPTY: (
            "{name} carries no text, so there is nothing to index.",
            "A scanned page holds a picture of text and not the text; upload a document whose "
            "words can be selected.",
        ),
    }
)
"""Code -> (message template, suggestion) for every failure this module reports (plan section 13.4).

The same shape as `engine.stages.score.SCORE_ERRORS`, and for the same reason: plan section 13.4
asks every user-facing error to carry a code, a message and a suggestion, and the API can only map
a code onto an HTTP status for codes that are written down. Every raise goes through
:func:`parse_error`, so a code that is not in this table cannot be raised.

A message names the document the caller uploaded, because that is what tells them which of their
own files to fix. It never reaches a log: the log line carries the code and the media type, and a
filename is customer data the moment it leaves the caller's own screen (plan section 13.7).
"""


class ParseError(Exception):
    """A document cannot be read into sections.

    `code` is machine-readable and `message` is business language, the same shape as `ScoreError`
    and `TrainError`, so `engine.errors.run_error` turns any of them into the same `RunError`
    without this module having to know that it exists. `suggestion` comes from
    :data:`PARSE_ERRORS`; `RunError` has no place for it, so it reaches the API and the log line
    that reports the refusal rather than `status.json`.
    """

    def __init__(self, code: str, message: str, *, suggestion: str = "") -> None:
        super().__init__(f"{code}: {message}")
        self.code = code
        self.message = message
        self.suggestion = suggestion if suggestion else PARSE_ERRORS.get(code, ("", ""))[1]


def parse_error(code: str, **values: object) -> ParseError:
    """Build the `ParseError` for `code` from :data:`PARSE_ERRORS`, filling its placeholders."""
    message, suggestion = PARSE_ERRORS[code]
    return ParseError(code, message.format(**values), suggestion=suggestion)


# ---------------------------------------------------------------------------
# 3. What a parsed document is
# ---------------------------------------------------------------------------
@dataclass(frozen=True)
class Section:
    """One headed run of a document: the unit a chunk is cut from and a citation points at."""

    heading: str
    """The heading as written, or the document's own name when the document carries none."""

    text: str
    """The prose under the heading, paragraphs separated by a blank line; empty for a bare heading."""

    page: int | None
    """Page the heading sits on, counting from 1; `None` for a format that has no pages."""

    ordinal: int
    """Position of the section within the document, from 0, in reading order."""


@dataclass(frozen=True)
class ParsedDocument:
    """One document as the index builder sees it: nothing about storage, nothing about embedding."""

    name: str
    """Filename as uploaded, which is what a citation shows."""

    media_type: str
    """The `DocumentType` value the reader was chosen by: pdf, docx, md or txt."""

    sections: tuple[Section, ...]
    """Sections in reading order; never empty, because an empty document is a `ParseError`."""

    pages: int | None
    """Pages the document has, or `None` for a format that has none. Measured, never guessed."""

    warnings: tuple[str, ...]
    """Machine-readable warnings raised while reading, in the order they were raised."""


@dataclass(frozen=True)
class _Piece:
    """One run of text a reader recovered, before the shared section rules are applied to it."""

    text: str
    page: int | None
    is_heading: bool


@dataclass(frozen=True)
class _Reading:
    """What one reader got out of one file. `pages` is `None` for the three unpaged formats."""

    pieces: tuple[_Piece, ...]
    pages: int | None
    warnings: tuple[str, ...]


# ---------------------------------------------------------------------------
# 4. The shape a heading has in a format that does not mark them
# ---------------------------------------------------------------------------
def _heading_shaped(line: str, *, allow_separator: bool = False) -> bool:
    """True when `line` looks like a heading on its own, before its neighbours are considered.

    `allow_separator` is set only for the first line of a PDF. Everywhere else a line carrying
    :data:`CELL_SEPARATOR` is a table row, because that is how all four readers render one - and a
    PDF that folded an em dash into a hyphen produces exactly the same characters, so mid-document
    the two are indistinguishable and the safe reading is "row".
    """
    text = line.strip()
    if not text or len(text) > HEADING_MAX_CHARS:
        return False
    if text[-1] in _HEADING_FORBIDDEN_ENDINGS:
        return False
    if not allow_separator and CELL_SEPARATOR in text:
        return False
    if len(text.split()) > HEADING_MAX_WORDS:
        return False
    return text[0].isupper()


def _ends_a_sentence(line: str) -> bool:
    """True when `line` finished what it was saying, which is what precedes a heading."""
    text = line.strip()
    return not text or text.endswith(_SENTENCE_ENDINGS)


# ---------------------------------------------------------------------------
# 5. The readers, one per DocumentType
# ---------------------------------------------------------------------------
def _read_pdf(path: Path, name: str) -> _Reading:
    """Read a PDF page by page, recovering headings from the shape of the flattened lines.

    A page that will not yield text is recorded as :data:`DOCUMENT_PAGE_UNREADABLE` and skipped
    rather than failing the document: fourteen readable pages and one that is a photograph is a
    knowledge base with a gap in it, and refusing the whole file would be the larger loss.
    """
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
        page_count = len(reader.pages)
    except Exception as exc:
        log_failure(_LOGGER, "parse.pdf.open", exc)
        raise parse_error(DOCUMENT_UNREADABLE, name=name) from exc

    warnings: list[str] = []
    numbered: list[tuple[int, str]] = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            extracted = page.extract_text()
        except Exception as exc:
            log_failure(_LOGGER, "parse.pdf.page", exc)
            warnings.append(DOCUMENT_PAGE_UNREADABLE)
            continue
        for raw in extracted.splitlines():
            line = _WHITESPACE_RUN.sub(" ", raw).strip()
            if line:
                numbered.append((number, line))

    return _Reading(pieces=_pdf_pieces(numbered), pages=page_count, warnings=tuple(warnings))


def _pdf_pieces(numbered: Sequence[tuple[int, str]]) -> tuple[_Piece, ...]:
    """Split a PDF's lines into headings and body runs, consecutive body lines rejoined by a space.

    Rejoining is what undoes the hard wrapping: a sentence broken across two lines is a sentence a
    chunker would split and a retriever would only half match.
    """
    pieces: list[_Piece] = []
    body: list[str] = []
    body_page: int | None = None
    previous_was_heading = False

    for position, (number, line) in enumerate(numbered):
        if position == 0:
            is_heading = _heading_shaped(line, allow_separator=True)
        else:
            is_heading = (
                _heading_shaped(line)
                and (previous_was_heading or _ends_a_sentence(numbered[position - 1][1]))
                and position + 1 < len(numbered)
            )
        if is_heading:
            if body:
                pieces.append(_Piece(" ".join(body), body_page, False))
                body = []
            pieces.append(_Piece(line, number, True))
            previous_was_heading = True
            continue
        if not body:
            body_page = number
        body.append(line)
        previous_was_heading = False

    if body:
        pieces.append(_Piece(" ".join(body), body_page, False))
    return tuple(pieces)


def _read_docx(path: Path, name: str) -> _Reading:
    """Read a DOCX in document order: a paragraph styled `Heading…` opens a section, a table is rows.

    `iter_inner_content` is used rather than `paragraphs` and `tables` in turn, because those two
    lists lose the order the body was written in and a table would land after the prose that
    follows it.
    """
    from docx import Document
    from docx.table import Table

    try:
        document = Document(str(path))
        contents = list(document.iter_inner_content())
    except Exception as exc:
        log_failure(_LOGGER, "parse.docx.open", exc)
        raise parse_error(DOCUMENT_UNREADABLE, name=name) from exc

    pieces: list[_Piece] = []
    for content in contents:
        if isinstance(content, Table):
            for row in content.rows:
                cells = CELL_SEPARATOR.join(cell.text.strip() for cell in row.cells)
                if cells.strip(CELL_SEPARATOR):
                    pieces.append(_Piece(cells, None, False))
            continue
        text = content.text.strip()
        if not text:
            continue
        style = content.style
        style_name = style.name if style is not None and style.name is not None else ""
        pieces.append(_Piece(text, None, style_name.startswith(_DOCX_HEADING_STYLES)))
    return _Reading(pieces=tuple(pieces), pages=None, warnings=())


def _inline_text(token: Token) -> str:
    """The characters of one Markdown inline run: the words, without the markup around them.

    Walking the children rather than taking `token.content` is what drops the emphasis markers, so
    an italic note reads as a sentence and not as one wrapped in asterisks.
    """
    if not token.children:
        return token.content.strip()
    parts: list[str] = []
    for child in token.children:
        if child.type in _MD_TEXT_CHILDREN:
            parts.append(child.content)
        elif child.type in _MD_BREAK_CHILDREN:
            parts.append(" ")
    return "".join(parts).strip()


def _read_markdown(path: Path, name: str) -> _Reading:
    """Read Markdown through its own parser: an ATX heading opens a section, a table row is a line.

    Tables are enabled on top of CommonMark deliberately. They are the one construct in this corpus
    whose meaning is lost by reading the raw characters, and a tariff table read as pipes is a
    tariff nobody can quote.
    """
    from markdown_it import MarkdownIt

    try:
        source = path.read_text(encoding="utf-8")
    except (OSError, UnicodeError) as exc:
        log_failure(_LOGGER, "parse.md.read", exc)
        raise parse_error(DOCUMENT_UNREADABLE, name=name) from exc

    tokens = MarkdownIt("commonmark").enable("table").parse(source)
    pieces: list[_Piece] = []
    row: list[str] = []
    in_table = False
    in_heading = False

    for token in tokens:
        if token.type == "table_open":
            in_table = True
        elif token.type == "table_close":
            in_table = False
        elif token.type == "tr_open":
            row = []
        elif token.type == "tr_close":
            cells = CELL_SEPARATOR.join(row)
            if cells.strip(CELL_SEPARATOR):
                pieces.append(_Piece(cells, None, False))
        elif token.type == "heading_open":
            in_heading = True
        elif token.type == "heading_close":
            in_heading = False
        elif token.type == "inline":
            text = _inline_text(token)
            if in_table:
                row.append(text)
            elif text:
                pieces.append(_Piece(text, None, in_heading))
        elif token.type in _MD_BODY_TOKENS and token.content.strip():
            pieces.append(_Piece(token.content.strip(), None, False))

    return _Reading(pieces=tuple(pieces), pages=None, warnings=())


def _text_runs(lines: Sequence[str]) -> list[str]:
    """Group plain-text body lines into paragraphs, a table row standing on its own.

    A line of plain text is a whole logical line, so a line carrying :data:`CELL_SEPARATOR` is a
    table row and joining it to the prose around it would produce one unsplittable run of a
    paragraph and a tariff table. The PDF reader deliberately does not do the same: there a line is
    a wrap fragment, and a sentence with an em dash in it would be cut in two.
    """
    runs: list[str] = []
    paragraph: list[str] = []
    for line in lines:
        if CELL_SEPARATOR in line:
            if paragraph:
                runs.append(" ".join(paragraph))
                paragraph = []
            runs.append(line)
            continue
        paragraph.append(line)
    if paragraph:
        runs.append(" ".join(paragraph))
    return runs


def _text_block_pieces(block: Sequence[str]) -> tuple[_Piece, ...]:
    """Split one blank-line-separated block of plain text into its heading and its prose.

    The first line is a heading when it has the shape of one. So is the *last* line, when the block
    has more than one: a writer that renders a document to plain text puts the blank line after a
    heading and not before it, which leaves each heading stranded at the foot of the paragraph that
    precedes it. Reading only the first line would have found the title and nothing else.
    """
    head = block[0] if _heading_shaped(block[0]) else None
    tail = block[-1] if len(block) > 1 and _heading_shaped(block[-1]) else None
    middle = block[1 if head else 0 : -1 if tail else None]

    pieces: list[_Piece] = []
    if head:
        pieces.append(_Piece(head, None, True))
    pieces.extend(_Piece(run, None, False) for run in _text_runs(middle))
    if tail:
        pieces.append(_Piece(tail, None, True))
    return tuple(pieces)


def _read_text(path: Path, name: str) -> _Reading:
    """Read plain text: blocks separated by a blank line, headings recovered from their shape."""
    try:
        source = path.read_text(encoding="utf-8")
    except (OSError, UnicodeError) as exc:
        log_failure(_LOGGER, "parse.txt.read", exc)
        raise parse_error(DOCUMENT_UNREADABLE, name=name) from exc

    pieces: list[_Piece] = []
    block: list[str] = []
    for raw in [*source.splitlines(), ""]:
        line = raw.strip()
        if line:
            block.append(line)
            continue
        if block:
            pieces.extend(_text_block_pieces(block))
            block = []
    return _Reading(pieces=tuple(pieces), pages=None, warnings=())


DOCUMENT_PARSERS: Final[Mapping[DocumentType, Callable[[Path, str], _Reading]]] = MappingProxyType(
    {
        DocumentType.PDF: _read_pdf,
        DocumentType.DOCX: _read_docx,
        DocumentType.MD: _read_markdown,
        DocumentType.TXT: _read_text,
    }
)
"""`DocumentType` -> the reader that owns it. The one place a format is dispatched on.

Looked up rather than branched on, so adding a format is adding a row and a function. A
`DocumentType` with no row here is reported as `DOCUMENT_TYPE_UNSUPPORTED` rather than raising a
`KeyError`, because a configuration that accepts a type this table has not been taught is an
operator's mistake and deserves a sentence, not a stack trace.
"""


# ---------------------------------------------------------------------------
# 6. Assembling sections and the public entry point
# ---------------------------------------------------------------------------
def _assemble(pieces: Sequence[_Piece], name: str) -> tuple[Section, ...]:
    """Group a reader's pieces into sections, in reading order.

    Body text that arrives before any heading is given the document's own name as its heading, and
    that is also what a document with no heading at all ends up as: one section, named after the
    file. A heading with nothing under it is kept as an empty section rather than dropped, because
    "this document has a section called Notice and it is empty" is true and useful, and the chunker
    simply passes over it.
    """
    sections: list[Section] = []
    heading = name
    page: int | None = None
    body: list[str] = []
    opened = False

    def close() -> None:
        sections.append(Section(heading=heading, text="\n\n".join(body), page=page, ordinal=len(sections)))

    for piece in pieces:
        if piece.is_heading:
            if opened or body:
                close()
            heading, page, body, opened = piece.text, piece.page, [], True
            continue
        if not body and not opened:
            page = piece.page
        body.append(piece.text)
    if opened or body:
        close()
    return tuple(sections)


def parse(path: Path, *, name: str | None = None) -> ParsedDocument:
    """Read the document at `path` into ordered, headed sections.

    `name` is the filename as the caller knows it, which is not always the name on disk: an upload
    is stored under a generated key and the document a citation has to show is the one the user
    recognises. It defaults to the path's own name.

    Raises :class:`ParseError` with `DOCUMENT_TYPE_UNSUPPORTED`, `DOCUMENT_UNREADABLE` or
    `DOCUMENT_EMPTY`. Nothing else: a document that is merely badly written is a document.
    """
    document_name = name if name else path.name
    suffix = path.suffix.lower().removeprefix(".")
    try:
        document_type = DocumentType(suffix)
    except ValueError as exc:
        raise parse_error(
            DOCUMENT_TYPE_UNSUPPORTED, name=document_name, suffix=suffix if suffix else "extensionless"
        ) from exc

    reader = DOCUMENT_PARSERS.get(document_type)
    if reader is None:  # pragma: no cover - unreachable while the table covers the enum
        raise parse_error(DOCUMENT_TYPE_UNSUPPORTED, name=document_name, suffix=suffix)
    if not path.is_file():
        raise parse_error(DOCUMENT_UNREADABLE, name=document_name)

    reading = reader(path, document_name)
    sections = _assemble(reading.pieces, document_name)
    if not sections or not any(section.text.strip() for section in sections):
        raise parse_error(DOCUMENT_EMPTY, name=document_name)

    warnings = list(reading.warnings)
    if not any(piece.is_heading for piece in reading.pieces):
        warnings.append(DOCUMENT_NO_HEADINGS)
    _LOGGER.info(
        "parse.document media_type=%s sections=%d pages=%s warnings=%d",
        document_type.value,
        len(sections),
        "-" if reading.pages is None else reading.pages,
        len(warnings),
    )
    return ParsedDocument(
        name=document_name,
        media_type=document_type.value,
        sections=sections,
        pages=reading.pages,
        warnings=tuple(warnings),
    )


# ---------------------------------------------------------------------------
# 7. Personal data in free text
# ---------------------------------------------------------------------------
FREE_TEXT_DETECTORS: Final[tuple[PiiDetector, ...]] = tuple(
    detector
    for detector in PII_DETECTORS
    if detector.value_pattern is not None and detector.min_distinct_ratio == 0.0
)
"""The Phase 1 detectors whose value pattern can be trusted on its own, in `PII_DETECTORS` order.

`min_distinct_ratio` is `PiiDetector`'s own record of which shapes need a column's vocabulary
before they may be believed, and it is above zero for exactly one detector: `name`, whose pattern
is "one to four capitalised words". Over a column that is a roster of people; over a complaint it
is the first word of every sentence, and redacting on it would replace half the prose with markers
and leave nothing a reviewer could read. Selecting on that field rather than naming the detector
means the choice follows the detectors: a future shape that needs a column to be trusted is
excluded here the day it is added, without this module being edited.
"""


@dataclass(frozen=True)
class _Match:
    """Where one detector fired. The matched characters are deliberately not carried."""

    start: int
    end: int
    kind: str


def _pii_matches(text: str) -> tuple[_Match, ...]:
    """Every span a free-text detector matched, earliest first and longest first within a position.

    `finditer`, not `fullmatch`: the Phase 1 detectors are applied to a whole cell, and a complaint
    hides an e-mail address in the middle of a sentence.
    """
    found: list[_Match] = []
    for detector in FREE_TEXT_DETECTORS:
        pattern = detector.value_pattern
        if pattern is None:  # pragma: no cover - FREE_TEXT_DETECTORS already excludes these
            continue
        found.extend(_Match(m.start(), m.end(), detector.kind) for m in pattern.finditer(text))
    return tuple(sorted(found, key=lambda match: (match.start, -match.end)))


def _kinds_of(matches: Sequence[_Match]) -> tuple[str, ...]:
    """The kinds present in `matches`, in `FREE_TEXT_DETECTORS` order so the result is stable."""
    fired = {match.kind for match in matches}
    return tuple(detector.kind for detector in FREE_TEXT_DETECTORS if detector.kind in fired)


def find_pii(text: str) -> tuple[str, ...]:
    """Detector kinds present anywhere in `text`, in `FREE_TEXT_DETECTORS` order.

    Returns the kinds and never the values, exactly as `engine.stages.ingest.detect_pii` does: the
    caller wants to know that a document carries phone numbers, and telling it which ones would put
    them in whatever the caller logs next (plan section 13.7).
    """
    return _kinds_of(_pii_matches(text))


def redact(text: str) -> tuple[str, tuple[str, ...]]:
    """`text` with every detected identifier replaced by `REDACTED`, and the kinds that were found.

    Overlapping spans are merged before replacement - the phone and Aadhaar shapes both match a
    twelve-digit string, from different offsets - so a value can never be half covered. The kinds
    are every kind that matched, including one whose span another detector's span swallowed, so
    what this function reports and what :func:`find_pii` reports are the same list.
    """
    matches = _pii_matches(text)
    if not matches:
        return text, ()

    spans: list[tuple[int, int]] = []
    for match in matches:
        if spans and match.start <= spans[-1][1]:
            spans[-1] = (spans[-1][0], max(spans[-1][1], match.end))
        else:
            spans.append((match.start, match.end))

    parts: list[str] = []
    cursor = 0
    for start, end in spans:
        parts.append(text[cursor:start])
        parts.append(REDACTED)
        cursor = end
    parts.append(text[cursor:])
    return "".join(parts), _kinds_of(matches)
